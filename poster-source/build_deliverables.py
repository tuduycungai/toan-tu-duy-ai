from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.graphics.barcode import qr
from reportlab.graphics.shapes import Drawing
from reportlab.graphics import renderPDF
import html,fitz
ROOT=Path('/workspace/scratch/585c68fcc3d1');out=ROOT/'output/v6.2';out.mkdir(exist_ok=True)
TITLE='THIẾT KẾ WEBSITE TÍCH HỢP AI CHATBOT GỢI MỞ LỜI GIẢI THEO CÁC THAO TÁC TƯ DUY TOÁN HỌC TRONG GIẢI TOÁN TỨ GIÁC LỚP 8 THEO CHƯƠNG TRÌNH GIÁO DỤC PHỔ THÔNG 2018'
SHEET='https://docs.google.com/spreadsheets/d/1SWeS4Y1Q5IciFAv2Ud2rOHIHhZqLXWIdzjhryotAy-c/edit?usp=drivesdk'
WEB='https://tuduyhinhtoan8.netlify.app/'
parts=[
('Mở đầu và câu hỏi nghiên cứu','Phan Trí Dũng','Khoảng 1 phút 30 giây. Chỉ vào tiêu đề và ô phía trên bên trái.',[
'Kính thưa quý thầy cô và Ban giám khảo, em là Phan Trí Dũng. Cùng thực hiện dự án với em là bạn Phạm Nam Hải. Chúng em là học sinh lớp 12, Trường THPT Gia Định, dưới sự hướng dẫn của thầy Trần Ngọc Phú Sang.',
'Hôm nay, chúng em xin trình bày đề tài: '+TITLE.capitalize().replace('ai chatbot','AI chatbot')+'.',
'Khi giải một bài toán hình học, việc nhớ tính chất chưa chắc giúp học sinh biết bắt đầu từ đâu. Một hình vẽ có thể gợi ra dự đoán, nhưng để chứng minh, người học phải xác định giả thiết, chọn căn cứ và nối các bước lập luận.',
'Từ vấn đề đó, chúng em đặt câu hỏi: Làm sao để mỗi lượt chat giúp học sinh tự tiến một bước suy luận? Mục đích của sản phẩm là hỗ trợ học sinh lớp 8 đọc đúng giả thiết, giải thích được căn cứ và vận dụng cách nghĩ vào bài mới.',
'Poster được chia thành bốn phần theo Phụ lục 3. Sau phần câu hỏi và mục đích nghiên cứu, chúng em sẽ trình bày phương pháp, dữ liệu và phân tích, rồi giải thích kết quả cùng tính mới của sản phẩm.'
]),
('Phương pháp và cách học với sản phẩm','Phan Trí Dũng','Khoảng 2 phút. Chỉ vào ô phương pháp phía dưới bên trái.',[
'Chúng em tổ chức công việc theo ba bước. Trước hết là nghiên cứu tài liệu để lựa chọn các thao tác tư duy và kiểm tra điều kiện toán học. Tiếp theo là thiết kế nguyên mẫu, chuẩn hóa câu hỏi gợi mở, sơ đồ và học liệu. Cuối cùng là kiểm thử chức năng và tổng hợp phản hồi trải nghiệm.',
'Các thao tác trong sản phẩm gồm phân tích, tổng hợp, so sánh, tương tự hóa, khái quát hóa và đặc biệt hóa. Chẳng hạn, phân tích giúp học sinh tách điều đã cho và điều cần chứng minh. Tổng hợp giúp nối những quan hệ đã có thành một lập luận. Tương tự hóa giúp học sinh thử một cách suy nghĩ đã dùng ở tình huống gần giống.',
'Trong một lượt học, website hiện bốn phần: giả thiết, mục tiêu, kết nối kiến thức và bước học sinh tự làm. Chatbot đưa ra một câu hỏi vừa sức để học sinh trả lời kèm căn cứ. Khi học sinh còn vướng, hướng thiết kế là thu nhỏ nhiệm vụ và điều chỉnh gợi ý theo câu trả lời thực tế.',
'Sản phẩm trên poster có 5 chế độ học, 87 bài tập, 93 hình từ phiếu và 18 tình huống sư phạm. Hệ thống dựng hình sử dụng quy trình Gemini tạo mã TikZ, sau đó biên dịch thành hình SVG. Hình giúp quan sát và kiểm tra cấu hình, còn kết luận toán học vẫn cần lập luận.',
'Sau đây, bạn Hải sẽ minh họa một tình huống học tập và trình bày phần minh chứng.'
]),
('Minh họa quá trình suy luận','Phạm Nam Hải','Khoảng 1 phút 30 giây. Chỉ vào ví dụ trung điểm trong ô phương pháp.',[
'Em xin tiếp nối bằng bài toán: Cho tứ giác lồi ABCD. Các điểm E, F, G, H lần lượt là trung điểm của AB, BC, CD, DA. Hỏi EFGH là hình gì?',
'Khi nhìn hình, học sinh có thể dự đoán EFGH là hình bình hành. Để kiểm chứng, chatbot gợi hỏi: Trong tam giác ABC, đoạn EF có vai trò gì? Học sinh nhận ra EF là đường trung bình, nên EF song song với AC và EF bằng một nửa AC.',
'Tiếp theo, học sinh xét tương tự trong tam giác ADC. Đoạn HG cũng là đường trung bình, nên HG song song với AC và HG bằng một nửa AC. Từ đó, EF song song và bằng HG. Một tứ giác có một cặp cạnh đối vừa song song vừa bằng nhau là hình bình hành. Vì vậy, EFGH là hình bình hành.',
'Ở đây, người học đã thực hiện phân tích, tương tự hóa và tổng hợp. Với câu hỏi mở rộng, nếu AC vuông góc BD thì EFGH là hình chữ nhật; nếu AC bằng BD thì EFGH là hình thoi; nếu đồng thời có cả hai điều kiện thì EFGH là hình vuông. Những kết luận mở rộng này cũng phải dựa trên quan hệ đường trung bình, chứ không chỉ dựa vào hình vẽ.'
]),
('Dữ liệu và phân tích kết quả','Phạm Nam Hải','Khoảng 1 phút 30 giây. Chỉ vào biểu đồ và QR minh chứng phía trên bên phải.',[
'Phần minh chứng trải nghiệm tổng hợp 111 lượt phản hồi ngày 16 tháng 9 năm 2026, theo thang đánh giá từ 1 đến 5. Có 55 lượt chọn mức 5, 39 lượt chọn mức 4 và 17 lượt chọn mức 3. Không có lượt chọn mức 1 hoặc mức 2.',
'Như vậy, có 94 trên 111 lượt chọn mức 4 hoặc 5, tương ứng khoảng 84,7 phần trăm. Điểm trung bình là 4,34 trên 5. Các con số này được trình bày bằng biểu đồ để người xem đối chiếu được số lượt ở từng mức. Quý thầy cô có thể quét QR để xem bảng tính minh chứng.',
'Chúng em dùng đơn vị lượt phản hồi, không đồng nhất với số học sinh độc lập. Kết quả cho thấy phản hồi trải nghiệm nhìn chung tích cực trong dữ liệu đã thu thập. Tuy nhiên, khảo sát này chưa đủ để kết luận học sinh tiến bộ về điểm số hoặc năng lực chứng minh.',
'Ngoài khảo sát, hồ sơ gói nộp V6.1 ghi nhận 401 ca kiểm thử đạt và 1 ca bỏ qua. Kiểm thử kỹ thuật và khảo sát trải nghiệm trả lời hai câu hỏi khác nhau, vì vậy chúng em không dùng số ca kiểm thử để suy ra hiệu quả học tập. Những kết quả này thuộc hồ sơ V6.1, chưa phải kết quả kiểm chứng hiệu quả của bản sửa V6.2.'
]),
('Tính mới, giới hạn và kết luận','Phan Trí Dũng','Khoảng 1 phút 30 giây. Chỉ vào ô phía dưới bên phải rồi hướng về Ban giám khảo.',[
'Tính mới mà nhóm tập trung thể hiện gồm bốn điểm: chuyển thao tác tư duy thành nhiệm vụ cụ thể; đặt sơ đồ bốn trường ngay trong hội thoại; điều chỉnh gợi ý theo câu trả lời thực tế; kết hợp câu hỏi gợi mở với hình TikZ để học sinh đối chiếu giả thiết.',
'Giá trị mà chúng em hướng tới là hỗ trợ tự học và thảo luận có căn cứ. Mỗi lần trả lời, học sinh cần biết mình dùng điều kiện nào và vì sao điều kiện đó giúp đi đến kết luận.',
'Nhóm cũng nhận thấy giới hạn của sản phẩm. AI có thể hiểu chưa đúng đề hoặc sinh hình chưa phù hợp. Trên điện thoại, việc dựng hình nặng và mất trạng thái khi tải lại có thể ảnh hưởng trải nghiệm. Bản cập nhật kỹ thuật cần được kiểm thử trên thiết bị thực tế trước khi khẳng định mức độ ổn định.',
'Trong bước tiếp theo, chúng em dự kiến đánh giá trước và sau khi sử dụng, với tiêu chí rõ ràng về đọc giả thiết, lựa chọn căn cứ và trình bày chứng minh. Qua đó, nhóm mới có cơ sở đánh giá hiệu quả học tập.',
'Chúng em xin kết thúc phần báo cáo với thông điệp trên poster: Hiểu vì sao, rồi mới kết luận. Chúng em cảm ơn quý thầy cô và Ban giám khảo đã lắng nghe, và xin tiếp nhận các câu hỏi góp ý.'
])]
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(12);doc.styles['Normal'].paragraph_format.line_spacing=1.12;doc.styles['Normal'].paragraph_format.space_after=Pt(7)
for i,(head,speaker,cue,paras) in enumerate(parts):
 if i:doc.add_page_break()
 else:
  doc.add_paragraph('Bài thuyết trình poster khoa học',style='Title')
  doc.add_paragraph('Kịch bản hai học sinh • Thời lượng dự kiến 8–9 phút')
 doc.add_paragraph(str(i+1)+' '+head,style='Heading 1')
 p=doc.add_paragraph();p.add_run(speaker).bold=True;p.add_run(' — '+cue).italic=True
 for para in paras:doc.add_paragraph(para)
doc.add_page_break();doc.add_paragraph('Gợi ý chuẩn bị phần trao đổi',style='Heading 1')
qas=[('Điểm khác biệt với chatbot trả lời trực tiếp là gì?','Sản phẩm tổ chức lượt học theo giả thiết, mục tiêu, kết nối kiến thức và bước học sinh tự làm. Trọng tâm là yêu cầu học sinh giải thích căn cứ và điều chỉnh câu gợi ý theo phần học sinh đang vướng.'),('84,7% có chứng minh sản phẩm giúp học tốt hơn không?','Chưa. Đây là tỷ lệ lượt phản hồi chọn mức 4–5 về trải nghiệm. Để đánh giá hiệu quả học tập cần thiết kế trước–sau, tiêu chí chấm thống nhất và cách xử lý dữ liệu rõ ràng.'),('111 phản hồi có tương ứng 111 học sinh không?','Không nên khẳng định như vậy. Bảng tổng hợp tính theo lượt gửi, chưa khử trùng người trả lời. Khi phân tích sâu hơn, nhóm cần xác định người tham gia độc lập và quy tắc loại dữ liệu không phù hợp.'),('Nếu AI dựng hình hoặc gợi ý sai thì sao?','Cần đối chiếu lại đề, giả thiết và điều kiện của định lý. Hình vẽ là minh họa. Giáo viên và học sinh kiểm tra lập luận trước khi sử dụng kết quả.'),('AI tham gia xây dựng sản phẩm như thế nào?','Khi báo cáo, nhóm trình bày đúng phần đã thực hiện và phần có AI hỗ trợ theo nhật ký câu lệnh, mã nguồn và kiểm thử. Không nhận phần AI tạo thành công việc tự thực hiện nếu chưa có kiểm chứng và công bố phù hợp.'),('Bản sửa trên điện thoại đã ổn định hoàn toàn chưa?','Chỉ nêu những kiểm tra thực tế đã hoàn thành. Phân biệt thử nghiệm mô phỏng với kiểm tra trên iPhone hoặc Android thật, và phân biệt mã đã sửa với phiên bản đã triển khai online.')]
for q,a in qas:doc.add_paragraph(q,style='Heading 2');doc.add_paragraph(a)
doc.add_paragraph('Chuẩn bị trước khi báo cáo',style='Heading 2');doc.add_paragraph('Mở sẵn website và bảng tính minh chứng. Phân vai như trên là gợi ý trình bày, không thay cho phân công nghiên cứu thực tế. Đối chiếu tên đề tài, mã số và phường hoặc xã theo hồ sơ chính thức trước khi nộp.')
doc.save(out/'Loi_thuyet_trinh_poster.docx')
# Poster follows the appendix: metadata, then four panels in a two-by-two grid.
for n,f in [('R','DejaVuSans.ttf'),('B','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(n,'/usr/share/fonts/truetype/dejavu/'+f))
W,H=1684,1191;c=canvas.Canvas(str(out/'Poster_online_Phu_luc_3.pdf'),pagesize=(W,H));c.setTitle(TITLE)
navy='#0A2744';teal='#008A86';blue='#245BE0';gold='#FFBE2E'
def rect(x,y,w,h,color):c.setFillColor(HexColor(color));c.rect(x,H-y-h,w,h,fill=1,stroke=0)
def t(x,y,s,size=18,col=navy,bold=False):c.setFillColor(HexColor(col));c.setFont('B' if bold else 'R',size);c.drawString(x,H-y,s)
def para(x,y,w,s,size=19,col=navy,bold=False):
 p=Paragraph(s,ParagraphStyle('p',fontName='B' if bold else 'R',fontSize=size,leading=size*1.35,textColor=HexColor(col)));_,h=p.wrap(w,1000);p.drawOn(c,x,H-y-h);return h
rect(0,0,W,H,'#F1F7FF');rect(0,0,W,245,navy);t(35,35,'PHỤ LỤC 3   /   POSTER ONLINE',17,gold,True);t(1350,35,'MÃ SỐ: [BỔ SUNG]',17,gold,True)
para(35,54,W-70,TITLE,27,'#FFFFFF',True)
t(35,194,'Phan Trí Dũng (lớp 12, nhóm trưởng)  •  Phạm Nam Hải (lớp 12)  •  Trường THPT Gia Định',20,'#FFFFFF',True)
t(35,226,'GV hướng dẫn: Trần Ngọc Phú Sang  •  SĐT: 0903397503  •  Phường (xã): [BỔ SUNG THEO HỒ SƠ]',17,'#D1E6F7')
heads=['CÂU HỎI NGHIÊN CỨU – MỤC ĐÍCH NGHIÊN CỨU','DỮ LIỆU VÀ PHÂN TÍCH DỮ LIỆU','PHƯƠNG PHÁP NGHIÊN CỨU','GIẢI THÍCH – KẾT LUẬN – TÍNH MỚI CỦA ĐỀ TÀI']
coords=[(30,265),(857,265),(30,708),(857,708)]
for (x,y),head in zip(coords,heads):
 rect(x,y,797,422,'#FFFFFF');rect(x,y,797,50,teal if x==30 else blue);para(x+18,y+12,761,head,19,'#FFFFFF',True)
para(53,335,747,'<b>Vấn đề:</b> Học sinh biết tính chất nhưng khó chọn căn cứ, dễ suy diễn từ hình vẽ và chưa tự nối được các bước chứng minh.',21)
para(53,429,747,'<b>Câu hỏi:</b> Làm sao để mỗi lượt chat giúp học sinh tự tiến một bước suy luận?',26,teal,True)
para(53,528,747,'<b>Mục đích:</b> Đọc đúng giả thiết, giải thích được căn cứ, tự vận dụng vào bài mới.<br/><br/><b>Đối tượng:</b> Học sinh lớp 8 học chủ đề Tứ giác theo Chương trình GDPT 2018.',21)
t(880,375,'84,7%',49,teal,True);t(1190,375,'4,34 / 5',45,navy,True)
t(880,409,'94/111 lượt chọn mức 4–5',18);t(1190,409,'Điểm trung bình trải nghiệm',18)
t(880,445,'111 lượt phản hồi • 16/09/2026 • Thang 1–5',17)
for i,(n,v) in enumerate([(5,55),(4,39),(3,17),(2,0),(1,0)]):
 y=477+i*32;t(883,y+18,str(n),18,navy,True);rect(925,y,max(v*5,1),21,teal if n==5 else blue);t(1230,y+18,f'{v} lượt ({v/111*100:.1f}%)'.replace('.',','),17)
para(880,652,742,'Phản hồi trải nghiệm tích cực; chưa đủ kết luận hiệu quả học tập.',17,teal,True)
para(53,773,747,'<b>1.</b> Nghiên cứu tài liệu và điều kiện toán học.<br/><b>2.</b> Thiết kế nguyên mẫu, câu hỏi và học liệu.<br/><b>3.</b> Kiểm thử chức năng, tổng hợp phản hồi.',19)
para(53,870,747,'<b>Thao tác:</b> Phân tích, tổng hợp, so sánh, tương tự hóa, khái quát hóa, đặc biệt hóa.',19)
para(53,938,747,'<b>Một lượt học:</b> Giả thiết → Mục tiêu → Kết nối kiến thức → Bước học sinh tự làm.<br/><br/><b>Ví dụ:</b> E, F, G, H là trung điểm các cạnh tứ giác lồi ABCD. EF và HG cùng song song AC, cùng bằng AC/2. Suy ra EFGH là hình bình hành.',19)
para(880,773,743,'<b>Điểm mới của sản phẩm</b><br/>• Chuyển thao tác tư duy thành nhiệm vụ cụ thể.<br/>• Hiện sơ đồ bốn trường trong hội thoại.<br/>• Điều chỉnh gợi ý theo câu trả lời thực tế.<br/>• Kết hợp gợi mở với hình TikZ.',20)
para(880,951,743,'<b>Giá trị:</b> Hỗ trợ tự học, thảo luận có căn cứ.<br/><b>Bước tiếp theo:</b> Đánh giá trước–sau để xác định hiệu quả học tập.',19)
def qr_at(url,x,y,size):
 q=qr.QrCodeWidget(url);a,b,d,e=q.getBounds();dr=Drawing(size,size,transform=[size/(d-a),0,0,size/(e-b),0,0]);dr.add(q);renderPDF.draw(dr,c,x,H-y-size);c.linkURL(url,(x,H-y-size,x+size,H-y),relative=0)
qr_at(SHEET,880,1020,90);qr_at(WEB,1280,1020,90);t(980,1060,'Bảng tính minh chứng',16,teal,True);t(1380,1060,'Website sản phẩm',16,blue,True)
t(35,1163,'Nguồn: Poster dự án; Khảo sát thông tin người dùng + Minh chứng.xlsx; hồ sơ NCKH 2627.',15)
c.save();d=fitz.open(out/'Poster_online_Phu_luc_3.pdf');d[0].get_pixmap(matrix=fitz.Matrix(.85,.85)).save(ROOT/'qa-v6.2/poster.png')
# Responsive version packaged at /poster.html in the application.
sections=[
(heads[0],'<p>Học sinh biết tính chất nhưng khó chọn căn cứ và nối các bước chứng minh.</p><h2>Làm sao để mỗi lượt chat giúp học sinh tự tiến một bước suy luận?</h2><p><b>Mục đích:</b> Đọc đúng giả thiết, giải thích được căn cứ, tự vận dụng vào bài mới.</p><p><b>Đối tượng:</b> Học sinh lớp 8 học chủ đề Tứ giác theo CT GDPT 2018.</p>'),
(heads[1],'<h2>84,7% <small>lượt chọn 4–5</small> &nbsp; 4,34/5</h2><p>111 lượt phản hồi ngày 16/09/2026.</p>'+''.join(f'<div class="bar"><b>Mức {n}</b><meter min="0" max="60" value="{v}">{v}</meter><span>{v} lượt</span></div>' for n,v in [(5,55),(4,39),(3,17),(2,0),(1,0)])+'<p>Phản hồi trải nghiệm tích cực; cần đánh giá trước–sau để xác định hiệu quả học tập.</p>'),
(heads[2],'<ol><li>Nghiên cứu tài liệu và điều kiện toán học.</li><li>Thiết kế nguyên mẫu, câu hỏi và học liệu.</li><li>Kiểm thử chức năng, tổng hợp phản hồi.</li></ol><p>Phân tích, tổng hợp, so sánh, tương tự hóa, khái quát hóa, đặc biệt hóa.</p><h2>Giả thiết → Mục tiêu → Kết nối kiến thức → Bước học sinh tự làm</h2><p>Ví dụ trung điểm: EF và HG cùng song song AC, cùng bằng AC/2. Suy ra EFGH là hình bình hành.</p>'),
(heads[3],'<ul><li>Chuyển thao tác tư duy thành nhiệm vụ cụ thể.</li><li>Hiện sơ đồ bốn trường trong hội thoại.</li><li>Điều chỉnh gợi ý theo câu trả lời thực tế.</li><li>Kết hợp gợi mở với hình TikZ.</li></ul><p>Hỗ trợ tự học và thảo luận có căn cứ. Cần kiểm tra hình, lời gợi ý và đánh giá hiệu quả học tập bằng dữ liệu phù hợp.</p>')]
page='<!doctype html><html lang="vi"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Poster online Phụ lục 3</title><style>body{margin:0;background:#eef5ff;color:#0a2744;font:18px/1.55 Arial,sans-serif}header{background:#0a2744;color:white;padding:30px}header h1{font-size:clamp(22px,2.5vw,36px);line-height:1.3}header b{color:#ffbe2e}main{padding:24px;display:grid;grid-template-columns:1fr 1fr;gap:20px}section{background:white;padding:24px;border-top:7px solid #008a86;border-radius:8px}h2{font-size:24px;color:#008a86}.label{font-weight:bold;color:#245be0}small{font-size:16px}footer{padding:24px}a{color:#245be0}nav{display:flex;flex-wrap:wrap;gap:24px}.bar{display:flex;align-items:center;gap:12px;margin:8px 0}meter{width:55%;height:24px}@media(max-width:700px){main{grid-template-columns:1fr;padding:12px}header{padding:20px}section{padding:18px}}@media print{body{font-size:13px}main{grid-template-columns:1fr 1fr}nav{display:none}}</style><header><b>PHỤ LỤC 3 • MẪU POSTER ONLINE</b><p>Mã số: [BỔ SUNG]</p><h1>'+TITLE+'</h1><p>Phan Trí Dũng (lớp 12, nhóm trưởng) • Phạm Nam Hải (lớp 12)<br>Trường THPT Gia Định • GV hướng dẫn: Trần Ngọc Phú Sang<br>SĐT: 0903397503 • Phường (xã): [BỔ SUNG THEO HỒ SƠ]</p></header><main>'+''.join('<section><div class="label">'+h+'</div>'+v+'</section>' for h,v in sections)+'</main><footer><nav><a href="'+WEB+'">Trải nghiệm website</a><a href="'+SHEET+'">Xem bảng tính minh chứng</a><a href="/poster-phu-luc-3.pdf">Tải poster PDF</a></nav><p>Nguồn: Poster dự án; Khảo sát thông tin người dùng + Minh chứng.xlsx; hồ sơ NCKH 2627.</p></footer></html>'
(out/'Poster_online_Phu_luc_3.html').write_text(page.replace('/poster-phu-luc-3.pdf','Poster_online_Phu_luc_3.pdf'))
(ROOT/'AI-TUTOR-TU-GIAC-PRO-V6_2/public/poster.html').write_text(page)
import shutil
shutil.copy2(out/'Poster_online_Phu_luc_3.pdf',ROOT/'AI-TUTOR-TU-GIAC-PRO-V6_2/public/poster-phu-luc-3.pdf')
print('Created DOCX, PDF and responsive HTML')
